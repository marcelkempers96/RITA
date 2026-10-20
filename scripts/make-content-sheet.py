# Builds docs/content-input/RITA-website-content-to-fill.xlsx from src/content (run: python3 scripts/make-content-sheet.py).
import json, glob, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

ROOT = '/home/user/RITA/src/content'
SITE = 'https://rita-woad.vercel.app'
def load(col):
    out = []
    for f in sorted(glob.glob(f'{ROOT}/{col}/*.json')):
        d = json.load(open(f)); d['_id'] = os.path.basename(f)[:-5]; out.append(d)
    return out
PROG = {'IND': '01 Green industrialisation & policy', 'MIN': '02 Transition minerals & value chains', 'EN': '03 Energy & industry', 'INN': '04 Innovation, firms & work'}
TYPE = {'journal': 'Journal article', 'report': 'Report', 'working-paper': 'Working paper', 'policy-paper': 'Policy paper', 'book-chapter': 'Book chapter', 'background-paper': 'Background paper',
        'op-ed': 'Op-ed', 'research-blog': 'Blog', 'policy-brief': 'Policy brief'}
authors = lambda a: ', '.join(a)

F = 'Arial'
HEAD = PatternFill('solid', fgColor='111110'); HFONT = Font(name=F, bold=True, color='FFFFFF', size=10)
FILL = PatternFill('solid', fgColor='FFF2B3')   # yellow: to fill in
EX = PatternFill('solid', fgColor='EDEDED')     # grey: example row
SEC = PatternFill('solid', fgColor='E9E1D3')    # sand: section rows
thin = Side(style='thin', color='D0CEC8'); BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical='top')

wb = Workbook()
def sheet(title, cols, rows, example=None, intro=None, freeze_cols=1, yn_cols=()):
    """cols: list of (header, width, is_input). rows: list of lists (input cells may be '' or a prefilled current value)."""
    ws = wb.create_sheet(title)
    r = 1
    if intro:
        ws.cell(r, 1, intro).font = Font(name=F, italic=True, size=10, color='45433F')
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(cols)); ws.row_dimensions[r].height = 32
        ws.cell(r, 1).alignment = WRAP; r += 1
    hdr = r
    for c, (h, w, inp) in enumerate(cols, 1):
        cell = ws.cell(r, c, h + (' ✎' if inp else '')); cell.font = HFONT; cell.fill = HEAD; cell.alignment = WRAP; cell.border = BORDER
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[r].height = 30; r += 1
    if example:
        for c, v in enumerate(example, 1):
            cell = ws.cell(r, c, v); cell.fill = EX; cell.font = Font(name=F, size=9, italic=True, color='45433F'); cell.alignment = WRAP; cell.border = BORDER
        r += 1
    first = r
    for row in rows:
        if isinstance(row, str):  # section row
            cell = ws.cell(r, 1, row); cell.font = Font(name=F, bold=True, size=10); 
            for c in range(1, len(cols) + 1): ws.cell(r, c).fill = SEC
            r += 1; continue
        for c, v in enumerate(row, 1):
            cell = ws.cell(r, c, v); cell.font = Font(name=F, size=10); cell.alignment = WRAP; cell.border = BORDER
            if cols[c - 1][2]: cell.fill = FILL
            if isinstance(v, str) and v.startswith('http'):
                cell.hyperlink = v; cell.font = Font(name=F, size=10, color='1D6B55', underline='single')
        r += 1
    ws.freeze_panes = ws.cell(hdr + 1, freeze_cols + 1)
    ws.auto_filter.ref = f'A{hdr}:{get_column_letter(len(cols))}{max(r - 1, hdr)}'
    for col in yn_cols:
        dv = DataValidation(type='list', formula1='"Yes,No,Not sure"', allow_blank=True); ws.add_data_validation(dv)
        dv.add(f'{get_column_letter(col)}{first}:{get_column_letter(col)}{r - 1}')
    return ws, first, r - 1

# ---------- Start here ----------
ws = wb.active; ws.title = 'Start here'
ws.column_dimensions['A'].width = 30; ws.column_dimensions['B'].width = 95
lines = [
    ('RITA website: content to fill in', None),
    ('Website', SITE),
    ('Prepared', '20 October 2026, from the live site and its content files'),
    ('', None),
    ('How to use this file', None),
    ('1', 'Each tab is one kind of content. Fill in the yellow cells (columns marked ✎). Leave a cell empty if you don\'t know yet.'),
    ('2', 'The grey row at the top of each tab is an example of the format we need. You can overwrite or ignore it.'),
    ('3', 'Grey "current" columns show what the website says today. Correct them in the yellow column next to them only if they are wrong.'),
    ('4', 'Long texts (bios, Policy Insight full texts) can also be sent as Word files: write "see Word file" in the cell.'),
    ('5', 'Nothing on the website was made up: anything we did not have shows as a yellow "TBC" box on the site. This file lists every one of them.'),
    ('', None),
    ('Tabs', None),
    ('1 Institute & legal', 'Organisation details, domain and emails, legal pages, About and Fellowship texts, launch event.'),
    ('2 People', 'Gideon and Elvis profiles (bios, links, portraits) and co-authors to invite as associates.'),
    ('3 Publications', 'All publications in the research library: summaries, key findings, missing years, author order.'),
    ('4 Policy Insights', 'All 45 issues: date, LinkedIn link, approval to migrate, full text.'),
    ('5 Op-eds & blogs', 'Blogs, op-eds and external policy briefs: missing links and dates.'),
    ('6 Policy briefs', 'The 6 RITA launch briefs: authors, recommendations, texts.'),
    ('7 Projects', 'All projects: years, description, RITA\'s role.'),
    ('8 Partners & logos', 'Which organisations may be named, and logo permissions.'),
    ('9 Photos', 'Credits and permissions for the 25 photos now on the site.'),
    ('', None),
    ('Priority before launch', None),
    ('A', 'Tab 1: domain, legal entity, About story and mission approval.'),
    ('B', 'Tab 2: portraits, bios and profile links for Gideon and Elvis.'),
    ('C', 'Tab 3: author order for every co-authored publication (it is used in citations), and key findings for the flagship publications (marked "Flagship").'),
    ('D', 'Tab 4: approval to migrate the Policy Insights from LinkedIn, with dates.'),
]
for i, (a, b) in enumerate(lines, 1):
    ca = ws.cell(i, 1, a); cb = ws.cell(i, 2, b)
    for c in (ca, cb): c.font = Font(name=F, size=10); c.alignment = WRAP
    if b is None and a: ca.font = Font(name=F, bold=True, size=12 if i == 1 else 11)
ws.cell(1, 1).font = Font(name=F, bold=True, size=14)
ws.cell(2, 2).hyperlink = SITE; ws.cell(2, 2).font = Font(name=F, size=10, color='1D6B55', underline='single')
leg = len(lines) + 2
ws.cell(leg, 1, 'Legend').font = Font(name=F, bold=True, size=11)
ws.cell(leg + 1, 1, 'Yellow cell').fill = FILL; ws.cell(leg + 1, 2, 'Please fill in')
ws.cell(leg + 2, 1, 'Grey cell').fill = EX; ws.cell(leg + 2, 2, 'Example of the expected format')
ws.cell(leg + 3, 1, 'Sand row').fill = SEC; ws.cell(leg + 3, 2, 'Section heading')
for rr in range(leg + 1, leg + 4):
    for c in (1, 2): ws.cell(rr, c).font = Font(name=F, size=10)

# ---------- 1 Institute & legal ----------
inst = [
    'Organisation',
    ['Legal entity name', 'Footer, Privacy, Terms', 'TBC', ''],
    ['Registration number', 'Footer (© 2026 RITA · Registered no.)', 'TBC', ''],
    ['Registered address', 'Privacy, Terms', 'TBC', ''],
    ['Legal status (e.g. non-profit company, foundation)', 'About › facts', 'TBC', ''],
    ['Date founded', 'About › facts', 'TBC', ''],
    ['Bases', 'About, Contact', 'Johannesburg · Delft', ''],
    ['Governing law for Terms', 'Terms', 'TBC', ''],
    'Domain, emails and social media',
    ['Website domain', 'Everywhere', 'rita.africa (assumed)', ''],
    ['General enquiries email', 'Contact', 'info@rita.africa', ''],
    ['Media / press email', 'Contact, Media', 'press@rita.africa', ''],
    ['Policy & advisory email', 'Contact', 'policy@rita.africa', ''],
    ['Partnerships & funding email', 'Contact', 'partners@rita.africa', ''],
    ['Fellowship email', 'Fellowship, Contact', 'fellows@rita.africa', ''],
    ['Events email', 'Events', 'events@rita.africa', ''],
    ['Editor (submissions) email', 'Submit', 'editor@rita.africa', ''],
    ['Privacy email', 'Privacy', 'privacy@rita.africa', ''],
    ['Accessibility email', 'Accessibility', 'access@rita.africa', ''],
    ['LinkedIn page URL', 'Footer, Contact', 'TBC', ''],
    ['X (Twitter) handle', 'Footer, Contact', 'TBC', ''],
    ['YouTube channel URL', 'Footer, Contact', 'TBC', ''],
    'About page',
    ['Approve the mission statement', 'About', 'To produce rigorous, independent evidence on how African economies can industrialise through the green transition, and to put it in the hands of those who decide.', ''],
    ['Approve or edit "Our story"', 'About', 'RITA grew out of a long collaboration between Gideon Ndubuisi and Elvis Korku Avenyo, who met at UNU-MERIT and have worked together since 2018: on robotisation and exports, value chains and innovation, South Africa\'s electricity crises, gender and the green transition, and a 2025 World Bank report on regional mineral value chains. RITA makes that collaboration an institution that is Africa-led, open to fellows from across the continent, and focused on evidence that policymakers can use.', ''],
    ['Approve the values', 'About', 'Evidence first · Africa-led · Independent · Open (publish openly wherever rights allow) · Useful.', ''],
    'Partners & funding page',
    ['Approve the independence statement', 'About › Partners', 'RITA decides its research questions, methods and conclusions. Funders do not review findings before publication. We disclose every funder of every project, and authors declare conflicts of interest.', ''],
    ['How RITA is funded (annual breakdown)', 'About › Partners', 'TBC', ''],
    ['Funding policy (PDF)', 'About › Partners', 'Not available yet', ''],
    ['Annual report (PDF)', 'About › Partners', 'Not available yet', ''],
    'RITA Fellowship',
    ['Who can apply (eligibility)', 'Fellowship', 'African researchers within 10 years of their PhD [TBC]', ''],
    ['What fellows get (stipend, mentoring, support)', 'Fellowship', 'Stipend, mentoring by programme leads, publication support [TBC]', ''],
    ['When the call opens / closes', 'Fellowship', 'TBC', ''],
    ['FAQ: Can I apply from outside Africa?', 'Fellowship', 'TBC', ''],
    ['FAQ: Is the fellowship remote?', 'Fellowship', 'TBC', ''],
    ['FAQ: What languages can I work in?', 'Fellowship', 'English and French [confirm]', ''],
    'Opportunities',
    ['Call for contributions: who does it report to?', 'Opportunities', 'TBC', ''],
    ['Research Associate, Transition minerals: is this a real opening? (if yes, add details)', 'Opportunities', 'Not shown', ''],
    ['Any other open positions', 'Opportunities', 'None', ''],
    'Launch event',
    ['Date', 'Events', 'TBC', ''],
    ['Start time and timezone', 'Events', 'TBC', ''],
    ['Duration', 'Events', '60 min', ''],
    ['Moderator', 'Events', 'TBC', ''],
    ['Platform (Zoom, Teams…)', 'Events', 'Online (link emailed after registration)', ''],
    ['Other confirmed events (title, date, place)', 'Events', 'None yet (suggested: mineral value chains webinar; electricity & manufacturing dialogue in Johannesburg; solar PV brief launch)', ''],
    'Legal pages (to be reviewed by a lawyer for POPIA + GDPR)',
    ['Lawyer who will review Privacy, Cookies and Terms', 'Privacy, Cookies, Terms', 'TBC', ''],
    ['Content licence for RITA\'s own text', 'Terms', 'CC BY 4.0 unless stated [confirm]', ''],
    ['Known accessibility issues', 'Accessibility', 'TBC after testing', ''],
    'Media page',
    ['One-page fact sheet (PDF)', 'Media › press kit', 'Not available yet', ''],
    ['Brand guide (PDF)', 'Media › press kit', 'Not available yet', ''],
]
sheet('1 Institute & legal', [('Item', 42, False), ('Where on the site', 22, False), ('Current text on site', 60, False), ('Your input / correction', 60, True)],
      inst, example=['(example) Date founded', 'About › facts', 'TBC', '[Month Year]'], intro='Organisation-wide details. Fill in the yellow column; write "OK" if the current text is correct.')

# ---------- 2 People ----------
people = load('people')
pf = [('Short bio (about 60 words)', ''), ('Long bio (about 150 words): confirm or correct the current one', 'bio'), ('Role title', 'role'), ('Base (city)', 'base'),
      ('Email to show on profile', ''), ('LinkedIn URL', ''), ('Google Scholar URL', ''), ('ORCID iD', ''), ('Personal website', 'site'),
      ('Focus tags (5 keywords)', 'tags'), ('Topics for media comment', 'exp'), ('Awards (confirm)', 'awards'), ('Portrait photo sent? (4:5, at least 1600×2000 px, with photo release)', '')]
prow = []
for p in people:
    prow.append(p['name'])
    for label, key in pf:
        cur = {'bio': p.get('bio', ''), 'role': p.get('role', ''), 'base': p.get('base', ''), 'site': p.get('links', {}).get('site', ''),
               'tags': ' · '.join(p.get('focusTags', [])), 'exp': ' · '.join(p.get('expertise', [])), 'awards': ' · '.join(p.get('awards', []))}.get(key, '') or 'TBC'
        prow.append([p['name'], label, cur, ''])
prow.append('Fellows, associates and staff to list (only with their written consent)')
for co in ['Solomon Owusu', 'Woubet Kassa', 'Emmanuel B. Mensah', 'Rexford Asiamah', 'Christian Urom', 'Maty Konte', 'Fiona Tregenna', 'Antonio Andreoni', 'Shingie Chisoro', 'Jason Bell', 'Julius Nyamwena', 'Keyi Tang', 'Xiaolan Fu', 'Erika Kraemer-Mbula']:
    prow.append([co, 'Invite as associate? (Yes / No) + role, email', 'Frequent co-author (not listed on the site)', ''])
prow.append(['(Add any other person)', 'Name, group (fellow / associate / staff), role, base, email', '', ''])
sheet('2 People', [('Person', 22, False), ('Item', 46, False), ('Current text on site', 70, False), ('Your input / correction', 60, True)], prow,
      example=['(example)', 'ORCID iD', 'TBC', '0000-0000-0000-0000'], intro='Profile details for the programme leads, and people to invite. Portraits: send separately (JPG, natural light, plain background).')

# ---------- 3 Publications ----------
pubs = load('publications')
progs = {p['_id']: p for p in load('programmes')}
flag = set(sum((p['featured'] for p in progs.values()), []))
pubs.sort(key=lambda p: (p.get('archive', False), -(p.get('year') or 0), p['title']))
cols = [('No.', 5, False), ('Title', 48, False), ('Outlet', 30, False), ('Year', 7, False), ('Authors (order as on site)', 26, False), ('Type', 14, False), ('Programme', 22, False),
        ('Priority', 10, False), ('Link', 30, False), ('Page on website', 30, False),
        ('Year, if missing', 10, True), ('Author order correct?', 11, True), ('Correct author order', 24, True),
        ('In brief (2–3 plain sentences)', 45, True), ('Why it matters (2 short paragraphs for policymakers)', 45, True),
        ('Key finding 1', 32, True), ('Key finding 2', 32, True), ('Key finding 3', 32, True), ('Data & methods (1 paragraph)', 40, True),
        ('May we host the PDF? (Yes/No)', 11, True), ('Region (confirm)', 18, True), ('Other corrections', 30, True)]
rows = []
for i, p in enumerate(pubs, 1):
    pr = 'Flagship' if (p.get('featured') or p['_id'] in flag) else ('Archive' if p.get('archive') else '')
    rows.append([i, p['title'], p['publisher'], p.get('year') or 'TBC', authors(p['authors']), TYPE.get(p['type'], p['type']), PROG.get((p.get('programmes') or [''])[0], 'Archive'),
                 pr, p.get('externalUrl', '') or 'TBC', f"{SITE}/research/publications/{p['_id']}",
                 '' , 'Yes' if p.get('authorOrderConfirmed') else '', '', p.get('summary', ''), '', '', '', '', '', '', ', '.join(p.get('regions', [])), ''])
ex = ['e.g.', '(example row)', '', '', '', '', '', '', '', '', '2024', 'No', '[Surname], [Surname], [Surname]',
      '[Two or three plain sentences: what the paper asks and finds]', '[Why this matters for policymakers]', '[One finding, one sentence, with a number if possible]', '[Finding 2]', '[Finding 3]', '[Data source, country, years, method]', 'No', 'Southern Africa', '']
sheet('3 Publications', cols, rows, example=ex, freeze_cols=2, yn_cols=(12, 20),
      intro=f'All {len(pubs)} publications in the research library. Start with rows marked "Flagship". "In brief" is prefilled where we already had text: correct it if needed. Archive items are hidden by default on the site.')

# ---------- 4 Policy Insights ----------
ins = sorted(load('insights'), key=lambda x: -x['number'])
cols = [('No.', 5, False), ('Title', 55, False), ('Series', 22, False), ('Part', 7, False), ('Programme', 22, False), ('Page on website', 32, False),
        ('Approve migration to website? (Yes/No)', 12, True), ('Publication date (YYYY-MM-DD)', 13, True), ('LinkedIn post URL', 32, True),
        ('Full text (paste, or "see Word file")', 60, True), ('Pull quote (one sentence, optional)', 35, True), ('Title correction', 30, True)]
rows = [[x['number'], x['title'], x.get('series', ''), f"{x['part']} of {x['of']}" if x.get('part') else '', PROG[x['programme']], f"{SITE}/insights/{x['_id']}",
         '', '', '', '', x.get('pullQuote', ''), ''] for x in ins]
sheet('4 Policy Insights', cols, rows, example=['e.g.', '(example row)', '', '', '', '', 'Yes', '2026-MM-DD', 'https://www.linkedin.com/posts/…', 'see Word file', '[One sentence from the issue]', ''],
      freeze_cols=2, yn_cols=(7,), intro='All 45 issues. The website shows titles only until the texts are migrated from LinkedIn. A date and LinkedIn link for each issue is the most useful start.')

# ---------- 5 Op-eds & blogs ----------
ops = sorted(load('opeds'), key=lambda o: o['date'], reverse=True)
cols = [('Title', 55, False), ('Outlet', 28, False), ('Type', 12, False), ('Date on site', 11, False), ('Authors', 24, False), ('Link on site', 32, False),
        ('Link (if missing or wrong)', 32, True), ('Correct month (YYYY-MM)', 12, True), ('May we republish the full text on the site? (Yes/No)', 14, True), ('One-line summary (confirm or write)', 45, True), ('Other corrections', 28, True)]
rows = []
for o in ops:
    rows.append([o['title'], o['outlet'], TYPE.get(o['kind'], o['kind']), o['date'], authors(o['authors']), o.get('outletUrl', '') or 'MISSING',
                 '', 'Please confirm' if o.get('note') else '', '', o.get('summary', ''), ''])
sheet('5 Op-eds & blogs', cols, rows, example=['(example row)', '', '', '', '', 'MISSING', 'https://…', 'YYYY-MM', 'No', '[One sentence: the argument]', ''],
      freeze_cols=1, yn_cols=(9,), intro='Blogs, op-eds and policy briefs published elsewhere. "MISSING" = no link yet. Summaries are shown on the site; correct them if needed.')

# ---------- 6 Policy briefs ----------
pubtitle = {p['_id']: p['title'] for p in pubs}; optitle = {o['_id']: o['title'] for o in ops}
briefs = sorted(load('briefs'), key=lambda b: b['number'])
cols = [('No.', 5, False), ('Title (working)', 40, False), ('Programme', 22, False), ('Based on', 45, False),
        ('Final title', 30, True), ('Authors', 22, True), ('Publication date', 12, True), ('Recommendation 1 + why', 40, True), ('Recommendation 2 + why', 40, True), ('Recommendation 3 + why', 40, True),
        ('The issue (1–2 paragraphs)', 45, True), ('What the evidence shows (2–3 paragraphs)', 45, True), ('PDF ready? (Yes/No)', 10, True)]
rows = [[b['number'], b['title'], PROG[b['programme']], ' | '.join(optitle.get(r.split(':', 1)[1], r) if r.startswith('opeds:') else pubtitle.get(r, r) for r in b['basedOn']),
         '', '', '', '', '', '', '', '', ''] for b in briefs]
sheet('6 Policy briefs', cols, rows, example=['e.g.', '(example row)', '', '', '[Final title]', '[Author names]', 'YYYY-MM', '[Recommendation in one line]. Why: [one line]', '[…]', '[…]', '[1–2 paragraphs]', '[2–3 paragraphs]', 'No'],
      freeze_cols=2, yn_cols=(13,), intro='RITA\'s six launch briefs. They show as "Forthcoming" on the site until recommendations are added.')

# ---------- 7 Projects ----------
prj = load('projects'); prj.sort(key=lambda p: (p['projectStatus'] != 'ongoing', p['title']))
pname = {p['_id']: p['name'] for p in people}
cols = [('Project', 42, False), ('Status', 10, False), ('Funder / commissioner', 30, False), ('Period on site', 12, False), ('RITA team', 26, False), ('Summary on site', 45, False), ('Page on website', 30, False),
        ('Start year', 9, True), ('End year', 9, True), ('Partners (comma-separated)', 32, True), ('About the project (1–2 paragraphs)', 50, True), ('RITA\'s role', 32, True), ('Summary correction', 35, True), ('Funding amount (optional)', 14, True)]
rows = [[p['title'], p['projectStatus'].capitalize(), p['funder'], p.get('period', '') or 'TBC', ', '.join([pname.get(t, t) for t in p.get('team', [])] + p.get('externalTeam', [])), p['summary'],
         f"{SITE}/research/projects/{p['_id']}", p.get('startYear', ''), p.get('endYear', ''), ', '.join(p.get('partners', [])), '', p.get('ritaRole', ''), '', p.get('budget', '')] for p in prj]
sheet('7 Projects', cols, rows, example=['(example row)', '', '', '', '', '', '', 'YYYY', 'YYYY', '[Partner 1], [Partner 2]', '[What the project studies and how]', '[e.g. co-lead, analysis]', '', ''], freeze_cols=1,
      intro='Includes work done before RITA was established. Prefilled yellow cells are what we already have: correct or complete them.')

# ---------- 8 Partners ----------
par = sorted(load('partners'), key=lambda p: p.get('order', 99))
cols = [('Organisation', 30, False), ('Relationship (as on site)', 45, False), ('Kind', 12, False), ('Shown on home page', 10, False),
        ('May we name them? (Yes/No)', 11, True), ('Logo permission granted? (Yes/No)', 11, True), ('Date of permission', 12, True), ('Contact person', 26, True), ('Website URL', 28, True), ('Relationship correction', 35, True)]
rows = [[p['name'], p['relationship'], p['kind'], 'Yes' if p.get('trackRecord') else '', '', '', '', '', p.get('url', ''), ''] for p in par]
sheet('8 Partners & logos', cols, rows, example=['(example row)', '', '', '', 'Yes', 'Yes', 'YYYY-MM-DD', '[Name, email]', 'https://…', ''],
      yn_cols=(5, 6), intro='Organisations are shown as names only until logo permission is granted. Send approved logo files (SVG preferred) separately.')

# ---------- 9 Photos ----------
ph = json.load(open('/home/user/RITA/scripts/photos.json'))['photos']
cols = [('Photo file', 22, False), ('Description', 60, False), ('Used for', 34, False),
        ('Photographer / credit line', 30, True), ('Licence or permission', 26, True), ('Location (city, country)', 22, True), ('OK to use? (Yes/No)', 10, True)]
rows = [[p['file'], p['alt'], ', '.join(p['shots']) or ('Cards: ' + ', '.join(PROG[c] for c in p['pools'])), '', '', '', ''] for p in ph]
rows.append(['portrait-gideon.jpg', 'Portrait of Gideon Ndubuisi', 'Profile, Fellows, Home', 'Still needed', '', '', ''])
rows.append(['portrait-elvis.jpg', 'Portrait of Elvis Korku Avenyo', 'Profile, Fellows, Home', 'Still needed', '', '', ''])
rows.append(['founders.jpg', 'Gideon and Elvis together (optional)', 'About › Our story', 'Still needed', '', '', ''])
sheet('9 Photos', cols, rows, example=['(example row)', '', '', 'Photo: [Name] / [Agency]', '[e.g. own photo; licence name]', '[City, Country]', 'Yes'], yn_cols=(7,),
      intro='The 25 photos you supplied are on the site. The credit line appears on each photo. Please confirm permission, especially for photos with visible brands (Coca-Cola, Kioo, KVOLT, MSC).')

for w in wb.worksheets[1:]:
    w.sheet_view.zoomScale = 90
out = '/home/user/RITA/docs/content-input/RITA-website-content-to-fill.xlsx'
wb.save(out)
print(out, [ (w.title, w.max_row) for w in wb.worksheets])
